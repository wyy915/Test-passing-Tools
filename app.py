from __future__ import annotations

import io
import json
import mimetypes
import posixpath
import re
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from email.header import decode_header
from email import policy
from email.parser import BytesParser
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, unquote, urlparse
from xml.etree import ElementTree as ET

from pypdf import PdfReader
from docx import Document


ROOT = Path(__file__).resolve().parent
PUBLIC_DIR = ROOT / "public"
DATA_DIR = ROOT / "data"
UPLOAD_DIR = ROOT / "uploads"
DB_PATH = DATA_DIR / "question_bank.json"
MEDIA_DIR = UPLOAD_DIR / "images"

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".pptx"}
LEADING_MARKERS = r"[▲△◆◇★☆●○■□▶▷▸▹►▻]+"
QUESTION_START = re.compile(
    rf"^\s*(?:{LEADING_MARKERS}\s*)*(?:第\s*)?(\d{{1,4}})\s*[\.．、\)）:：]"
    rf"(?:\s+(.+)|([^\d\s].*))$"
)
OPTION_START = re.compile(rf"^\s*(?:{LEADING_MARKERS}\s*)*([A-Fa-f])\s*[\.．、\)）:：]\s*(.*)$")
INLINE_OPTION_START = re.compile(rf"(?<![A-Za-z])(?:{LEADING_MARKERS}\s*)*([A-Fa-f])\s*[\.．:：]\s*")
EMBEDDED_QUESTION_START = re.compile(
    rf"(?<!\d)((?:{LEADING_MARKERS}\s*)*\d{{1,4}}\s*[\.．、\)）:：]"
    rf"\s+(?=\S)(?=\D))"
)
MARKER_ONLY_LINE = re.compile(rf"^\s*(?:{LEADING_MARKERS}\s*)+$")
TRAILING_MARKERS = re.compile(rf"(?:\s*{LEADING_MARKERS})+\s*$")
ANSWER_LINE = re.compile(
    r"^\s*(?:\d{1,4}\s*[\.．、\)）]\s*)?(?:参考)?(?:【\s*)?答案(?:\s*】)?(?:为|是)?\s*[:：]?\s*(.+?)\s*$",
    re.IGNORECASE,
)
EXPLANATION_LINE = re.compile(
    r"^\s*(?:\d{1,4}\s*[\.．、\)）]\s*)?(?:答案)?(?:【\s*)?解析(?:\s*】)?\s*[:：]?\s*(.*)$",
    re.IGNORECASE,
)
ANSWER_SECTION_MARKER = re.compile(r"(?:答案与解析|答案解析|参考答案)", re.IGNORECASE)
ANSWER_ITEM_LINE = re.compile(
    rf"^\s*(?:{LEADING_MARKERS}\s*)*(\d{{1,4}})\s*[\.．、\)）:：]?\s*"
    r"(?:参考)?(?:【\s*)?答案(?:\s*】)?(?:为|是)?\s*[:：。．.!！]?\s*"
    r"([A-Fa-f]{1,6}|正确|错误|对|错)\s*[。．.!！]?\s*$",
    re.IGNORECASE,
)
ANSWER_KEY_ITEM = re.compile(
    rf"(?<!\d)(?:{LEADING_MARKERS}\s*)*(\d{{1,4}})\s*[\.．、\)）:：]?\s*"
    r"(?:参考)?(?:【\s*)?答案(?:\s*】)?(?:为|是)?\s*[:：]?\s*"
    r"([A-Fa-f]{1,6}|正确|错误|对|错)(?=\s|$|[,，;；])",
    re.IGNORECASE,
)
PLAIN_ANSWER_KEY_ITEM = re.compile(
    rf"(?<!\d)(?:{LEADING_MARKERS}\s*)*(\d{{1,4}})\s*[\.．、\)）:：]?\s*"
    r"([A-Fa-f]{1,6}|正确|错误|对|错)(?=\s|$|[,，;；])",
    re.IGNORECASE,
)
SECTION_HEADING = re.compile(
    r"^\s*(?:[一二三四五六七八九十百千万\d]+[、.．:：)\]]\s*)?"
    r"(单项选择题|单选题|多项选择题|多选题|判断题|填空题|简答题|选择题)"
    r"(?:\s*[\(（].*[\)）])?\s*$",
    re.IGNORECASE,
)


@dataclass
class AnswerRecord:
    number: str
    answer: list[str]
    explanation: str = ""
    section: str | None = None


@dataclass
class PreparedAnswerData:
    answers: dict[str, list[str]]
    explanations: dict[str, str]
    records: list[AnswerRecord]

    def __iter__(self):
        yield self.answers
        yield self.explanations

    def __getitem__(self, index: int):
        if index == 0:
            return self.answers
        if index == 1:
            return self.explanations
        raise IndexError(index)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def ensure_storage() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    UPLOAD_DIR.mkdir(exist_ok=True)
    MEDIA_DIR.mkdir(exist_ok=True)
    if not DB_PATH.exists():
        DB_PATH.write_text(
            json.dumps({"modules": [], "wrongBook": [], "practiceProgress": {}}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )


def load_db() -> dict[str, Any]:
    ensure_storage()
    try:
        payload = json.loads(DB_PATH.read_text(encoding="utf-8"))
        if isinstance(payload, dict) and isinstance(payload.get("modules"), list):
            if not isinstance(payload.get("wrongBook"), list):
                payload["wrongBook"] = []
            if not isinstance(payload.get("practiceProgress"), dict):
                payload["practiceProgress"] = {}
            for entry in payload["wrongBook"]:
                entry["category"] = normalize_category(entry.get("category", "未分类"))
                entry.setdefault("note", "")
                entry.setdefault("manualAnswer", [])
            for module in payload["modules"]:
                module.setdefault("questions", [])
                module.setdefault("knowledgePoints", [])
                for question in module["questions"]:
                    question.setdefault("images", [])
            return payload
    except (OSError, json.JSONDecodeError):
        pass
    return {"modules": [], "wrongBook": [], "practiceProgress": {}}


def save_db(payload: dict[str, Any]) -> None:
    tmp_path = DB_PATH.with_suffix(".tmp")
    tmp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp_path.replace(DB_PATH)


def normalize_name(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = re.sub(r"\s+", " ", value).strip(" ._-\t\r\n")
    return value or "未命名题库"


def normalize_category(value: str) -> str:
    value = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value or ""))).strip()
    return value[:40] or "未分类"


def decode_filename(value: str) -> str:
    if not value:
        return "未命名文件"
    try:
        pieces = decode_header(value)
        return "".join(
            part.decode(charset or "utf-8", errors="replace") if isinstance(part, bytes) else part
            for part, charset in pieces
        )
    except Exception:
        return value


def read_uploaded_files(body: bytes, content_type: str) -> list[tuple[str, str, bytes]]:
    headers = (
        f"Content-Type: {content_type}\r\n"
        "MIME-Version: 1.0\r\n"
        "\r\n"
    ).encode("utf-8")
    message = BytesParser(policy=policy.default).parsebytes(headers + body)
    uploads: list[tuple[str, str, bytes]] = []
    for part in message.walk():
        if part.is_multipart():
            continue
        field_name = part.get_param("name", header="content-disposition") or ""
        filename = part.get_filename()
        if not filename or field_name not in {"file", "files", "question_file", "answer_file"}:
            continue
        content = part.get_payload(decode=True) or b""
        uploads.append((field_name, decode_filename(filename), content))
    return uploads


def read_uploaded_file(body: bytes, content_type: str) -> tuple[str, bytes] | None:
    """Backward-compatible single-file helper."""
    uploads = read_uploaded_files(body, content_type)
    if not uploads:
        return None
    _, filename, content = uploads[0]
    return filename, content


IMAGE_MARKER = re.compile(r"\[\[IMAGE:([^\]]+)\]\]")
REL_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def save_media_asset(content: bytes, filename: str, media_dir: Path = MEDIA_DIR) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg"}:
        suffix = ".bin"
    asset_name = f"{uuid.uuid4().hex}{suffix}"
    media_dir.mkdir(parents=True, exist_ok=True)
    (media_dir / asset_name).write_bytes(content)
    return f"/media/{asset_name}"


def append_image_marker(
    markers: list[str],
    related_part: Any,
    cache: dict[str, str],
    media_dir: Path | None,
) -> None:
    if media_dir is None or related_part is None:
        return
    cache_key = str(getattr(related_part, "partname", "")) or str(id(related_part))
    image_url = cache.get(cache_key)
    if image_url is None:
        part_name = str(getattr(related_part, "partname", "image.png"))
        image_url = save_media_asset(
            getattr(related_part, "blob", b""),
            part_name.rsplit("/", 1)[-1],
            media_dir,
        )
        cache[cache_key] = image_url
    if image_url not in markers:
        markers.append(image_url)


def read_docx_paragraph(
    paragraph: Any,
    document: Document,
    media_dir: Path | None,
    image_cache: dict[str, str],
) -> str:
    parts: list[str] = []
    for element in paragraph._p.iter():
        if element.tag.endswith("}t"):
            parts.append(element.text or "")
        elif element.tag.endswith("}tab"):
            parts.append(" ")
        elif element.tag.endswith("}br"):
            parts.append("\n")
        elif element.tag.endswith("}blip"):
            relationship_id = element.attrib.get(REL_EMBED)
            related_part = document.part.related_parts.get(relationship_id) if relationship_id else None
            image_markers: list[str] = []
            append_image_marker(image_markers, related_part, image_cache, media_dir)
            parts.extend(f"[[IMAGE:{url}]]" for url in image_markers)
    return "".join(parts).strip()


def read_docx_text(content: bytes, media_dir: Path | None = None) -> str:
    document = Document(io.BytesIO(content))
    image_cache: dict[str, str] = {}
    rows: list[str] = []
    for paragraph in document.paragraphs:
        text = read_docx_paragraph(paragraph, document, media_dir, image_cache)
        if text:
            rows.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [
                read_docx_paragraph(paragraph, document, media_dir, image_cache)
                for cell in row.cells
                for paragraph in cell.paragraphs
            ]
            if any(cells):
                rows.append("    ".join(cells))
    return "\n".join(rows)


def read_pdf_text(content: bytes, media_dir: Path | None = None) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if media_dir is not None:
            image_markers: list[str] = []
            try:
                page_images = page.images
            except Exception:
                # A text-only PDF should still import when optional image
                # dependencies are unavailable.
                page_images = []
            for image in page_images:
                image_url = save_media_asset(image.data, image.name, media_dir)
                if image_url not in image_markers:
                    image_markers.append(image_url)
            if image_markers:
                page_text = f"{page_text}\n" + "".join(
                    f"[[IMAGE:{image_url}]]\n" for image_url in image_markers
                )
        pages.append(page_text)
    return "\n".join(pages)


def read_pptx_text(content: bytes, media_dir: Path | None = None) -> str:
    slide_text: list[str] = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        slide_names = sorted(
            name for name in archive.namelist()
            if re.match(r"ppt/slides/slide\d+\.xml$", name)
        )
        for slide_name in slide_names:
            root = ET.fromstring(archive.read(slide_name))
            texts = [
                node.text or ""
                for node in root.iter()
                if node.tag.endswith("}t") and (node.text or "").strip()
            ]
            slide_parts = [" ".join(texts)] if texts else []
            if media_dir is not None:
                relationship_name = posixpath.join(
                    posixpath.dirname(slide_name),
                    "_rels",
                    posixpath.basename(slide_name) + ".rels",
                )
                relationships: dict[str, str] = {}
                if relationship_name in archive.namelist():
                    rel_root = ET.fromstring(archive.read(relationship_name))
                    for relationship in rel_root:
                        rel_id = relationship.attrib.get("Id")
                        target = relationship.attrib.get("Target")
                        if rel_id and target:
                            relationships[rel_id] = posixpath.normpath(
                                posixpath.join(posixpath.dirname(slide_name), target)
                            )
                image_urls: list[str] = []
                for node in root.iter():
                    if not node.tag.endswith("}blip"):
                        continue
                    target = relationships.get(node.attrib.get(REL_EMBED, ""))
                    if target and target in archive.namelist():
                        image_url = save_media_asset(
                            archive.read(target),
                            target.rsplit("/", 1)[-1],
                            media_dir,
                        )
                        if image_url not in image_urls:
                            image_urls.append(image_url)
                slide_parts.extend(f"[[IMAGE:{url}]]" for url in image_urls)
            if slide_parts:
                slide_text.append("\n".join(slide_parts))
    return "\n".join(slide_text)


def read_file_text(filename: str, content: bytes, media_dir: Path | None = None) -> str:
    extension = Path(filename).suffix.lower()
    if extension == ".docx":
        return read_docx_text(content, media_dir)
    if extension == ".pdf":
        return read_pdf_text(content, media_dir)
    if extension == ".pptx":
        return read_pptx_text(content, media_dir)
    return content.decode("utf-8-sig", errors="replace")


def answer_file_score(filename: str, raw_text: str) -> int:
    name = Path(filename).stem.casefold()
    score = 0
    if any(token in name for token in ["答案", "解析", "参考", "answer", "solution", "key"]):
        score += 8
    if ANSWER_SECTION_MARKER.search(raw_text):
        score += 5
    lines = [line.strip() for line in raw_text.replace("\r\n", "\n").split("\n") if line.strip()]
    answer_map, _ = extract_answer_data(lines)
    score += min(len(answer_map), 20)
    question_count = sum(1 for line in lines if is_question_start(line))
    if question_count >= 3:
        score -= min(question_count, 10)
    return score


def prepare_answer_data(raw_text: str) -> PreparedAnswerData:
    lines = [line.strip() for line in raw_text.replace("\r\n", "\n").split("\n") if line.strip()]
    answers, explanations = extract_answer_data(lines)
    return PreparedAnswerData(answers, explanations, extract_answer_records(lines))


def classify_uploads(
    uploads: list[tuple[str, str, bytes]],
) -> tuple[tuple[str, str, bytes], tuple[str, str, bytes] | None]:
    """Choose the question document and optional answer document."""
    explicit_question = next((item for item in uploads if item[0] == "question_file"), None)
    explicit_answer = next((item for item in uploads if item[0] == "answer_file"), None)
    if explicit_question:
        remaining = [item for item in uploads if item is not explicit_question]
        return explicit_question, explicit_answer or (remaining[0] if remaining else None)
    if len(uploads) == 1:
        return uploads[0], None
    if len(uploads) != 2:
        raise ValueError("一次最多上传两个文件：一个题目文件和一个答案文件")

    scored: list[tuple[int, tuple[str, str, bytes]]] = []
    for item in uploads:
        try:
            raw_text = read_file_text(item[1], item[2])
            score = answer_file_score(item[1], raw_text)
        except Exception:
            score = 0
        scored.append((score, item))
    scored.sort(key=lambda item: item[0], reverse=True)
    answer_score, answer_upload = scored[0]
    question_score, question_upload = scored[1]
    if answer_score <= question_score and answer_score < 8:
        # When filenames/content do not identify the roles, keep picker order:
        # the first file is the question file and the second is the answer file.
        question_upload, answer_upload = uploads[0], uploads[1]
    return question_upload, answer_upload


def clean_answer(value: str) -> list[str]:
    normalized = unicodedata.normalize("NFKC", value).strip()
    normalized = re.sub(r"[（(]多选[）)]", "", normalized, flags=re.IGNORECASE)
    normalized = normalized.replace("正确", "对").replace("错误", "错")
    if normalized in {"对", "错"}:
        return [normalized]
    letters = re.findall(r"[A-Fa-f]", normalized)
    if letters:
        return sorted(set(letter.upper() for letter in letters))
    if normalized:
        return [normalized]
    return []


def normalize_manual_answer(value: Any, options: list[dict[str, str]]) -> list[str]:
    """Normalize manually entered answers without rejecting free-form corrections."""
    raw_values = value if isinstance(value, list) else [value]
    answer: list[str] = []
    for raw in raw_values:
        text = unicodedata.normalize("NFKC", str(raw or "")).strip()
        if not text:
            continue
        items = clean_answer(text) if len(text) <= 12 else [text]
        for item in items:
            candidate = item.upper() if len(item) == 1 and item.isalpha() else item
            if candidate and candidate not in answer:
                answer.append(candidate)
    return answer


def normalize_manual_options(value: Any) -> list[dict[str, str]]:
    """Normalize edited options; incomplete blank rows are ignored."""
    if not isinstance(value, list):
        return []
    options: list[dict[str, str]] = []
    seen: set[str] = set()
    for raw_option in value:
        if not isinstance(raw_option, dict):
            continue
        raw_key = unicodedata.normalize("NFKC", str(raw_option.get("key") or "")).strip()
        text = unicodedata.normalize("NFKC", str(raw_option.get("text") or "")).strip()
        if not raw_key and not text:
            continue
        key = raw_key.upper() if len(raw_key) == 1 and raw_key.isalpha() else raw_key
        if not key:
            key = chr(ord("A") + len(options)) if len(options) < 26 else str(len(options) + 1)
        if key in seen:
            continue
        seen.add(key)
        options.append({"key": key, "text": text})
    return options


def question_type(question_text: str, answer: list[str], options: list[dict[str, str]]) -> str:
    if answer and answer[0] in {"对", "错"}:
        return "判断题"
    if len(answer) > 1 or "多选" in question_text:
        return "多选题"
    if options and {option["key"] for option in options} <= {"对", "错"}:
        return "判断题"
    if "判断题" in question_text and not options:
        return "判断题"
    if not options and any(k in question_text for k in ["正确", "错误", "对错"]):
        return "判断题"
    return "单选题"


def question_stem(match: re.Match[str]) -> str:
    return (match.group(2) or match.group(3) or "").strip()


def normalize_section(value: str) -> str:
    normalized = re.sub(r"\s+", "", value).casefold()
    if "多选" in normalized or "多项选择" in normalized:
        return "多选题"
    if "单选" in normalized or "单项选择" in normalized:
        return "单选题"
    if "判断" in normalized:
        return "判断题"
    if "填空" in normalized:
        return "填空题"
    if "简答" in normalized:
        return "简答题"
    return "选择题"


def section_heading(line: str) -> str | None:
    match = SECTION_HEADING.match(line)
    return normalize_section(match.group(1)) if match else None


def is_question_start(line: str) -> bool:
    return bool(QUESTION_START.match(line)) and not ANSWER_LINE.match(line) and not EXPLANATION_LINE.match(line)


def split_embedded_question_lines(lines: list[str]) -> list[str]:
    result: list[str] = []
    for line in lines:
        remainder = line
        while True:
            match = EMBEDDED_QUESTION_START.search(remainder, 1)
            if not match:
                break
            prefix = remainder[:match.start()].strip()
            if prefix:
                result.append(prefix)
            remainder = remainder[match.start():].strip()
        if remainder.strip():
            result.append(remainder.strip())
    return result


def extract_answer_records(lines: list[str]) -> list[AnswerRecord]:
    """Read answer records without letting the next question leak into an explanation."""
    records: list[AnswerRecord] = []
    current: AnswerRecord | None = None
    explanation_lines: list[str] = []
    current_section: str | None = None

    def flush() -> None:
        nonlocal current, explanation_lines
        if current is not None:
            current.explanation = " ".join(explanation_lines).strip()
            records.append(current)
        current = None
        explanation_lines = []

    for line in lines:
        heading = section_heading(line)
        if heading:
            flush()
            current_section = heading
            continue

        item_line = ANSWER_ITEM_LINE.match(line)
        if item_line:
            flush()
            current = AnswerRecord(
                number=item_line.group(1),
                answer=clean_answer(item_line.group(2)),
                section=current_section,
            )
            continue

        explanation_line = EXPLANATION_LINE.match(line)
        if explanation_line:
            number_match = re.match(r"^\s*(\d{1,4})\s*[\.．、\)）:：]", line)
            if number_match and (current is None or current.number != number_match.group(1)):
                flush()
                current = AnswerRecord(number=number_match.group(1), answer=[], section=current_section)
            explanation_text = explanation_line.group(1).strip()
            if current is not None and explanation_text:
                explanation_lines.append(explanation_text)
            continue

        key_matches = list(ANSWER_KEY_ITEM.finditer(line))
        if not key_matches:
            key_matches = list(PLAIN_ANSWER_KEY_ITEM.finditer(line))
        if key_matches:
            flush()
            for match in key_matches:
                records.append(
                    AnswerRecord(
                        number=match.group(1),
                        answer=clean_answer(match.group(2)),
                        section=current_section,
                    )
                )
            continue

        # A question-looking line starts a new boundary even when that question
        # has no answer entry in the answer document.
        if is_question_start(line):
            flush()
            continue

        if current is not None and line.strip() and not ANSWER_SECTION_MARKER.search(line):
            explanation_lines.append(line.strip())

    flush()
    return records


def extract_answer_data(lines: list[str]) -> tuple[dict[str, list[str]], dict[str, str]]:
    """Read a concentrated answer key and optional per-question explanations."""
    records = extract_answer_records(lines)
    answers = {record.number: record.answer for record in records}
    explanations = {
        record.number: record.explanation
        for record in records
        if record.explanation
    }
    return answers, explanations


class AnswerRecordMatcher:
    def __init__(self, sources: list[list[AnswerRecord]]) -> None:
        self.records = [record for source in sources for record in source]
        self.used: set[int] = set()

    def take(self, number: str, section: str | None) -> AnswerRecord | None:
        candidates: list[int] = []
        for index, record in enumerate(self.records):
            if index in self.used or record.number != number:
                continue
            if section and record.section == section:
                candidates.append(index)
        if not candidates:
            for index, record in enumerate(self.records):
                if index in self.used or record.number != number:
                    continue
                if record.section is None or section is None:
                    candidates.append(index)
        if not candidates:
            for index, record in enumerate(self.records):
                if index not in self.used and record.number == number:
                    candidates.append(index)
        if not candidates:
            return None
        selected = candidates[0]
        self.used.add(selected)
        return self.records[selected]


def parse_questions(
    raw_text: str,
    external_answer_data: PreparedAnswerData | tuple[dict[str, list[str]], dict[str, str]] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    lines = [line.strip() for line in text.split("\n")]
    lines = [line for line in lines if line and not MARKER_ONLY_LINE.match(line)]

    answer_section_index = next(
        (
            index
            for index, line in enumerate(lines)
            if ANSWER_SECTION_MARKER.search(line) and not is_question_start(line)
        ),
        len(lines),
    )
    question_lines = split_embedded_question_lines(lines[:answer_section_index])
    answer_lines = lines[answer_section_index:]
    embedded_records = extract_answer_records(answer_lines)
    answer_map, explanation_map = extract_answer_data(answer_lines)
    external_records: list[AnswerRecord] = []
    if external_answer_data:
        if isinstance(external_answer_data, PreparedAnswerData):
            external_answers = external_answer_data.answers
            external_explanations = external_answer_data.explanations
            external_records = external_answer_data.records
        else:
            external_answers, external_explanations = external_answer_data
            external_records = [
                AnswerRecord(number=number, answer=answer, explanation=external_explanations.get(number, ""))
                for number, answer in external_answers.items()
            ]
        answer_map.update(external_answers)
        explanation_map.update(external_explanations)
    answer_matcher = AnswerRecordMatcher([external_records, embedded_records])

    has_question_starts = any(is_question_start(line) for line in question_lines)
    if not has_question_starts:
        knowledge_points = extract_knowledge_points(
            [IMAGE_MARKER.sub("", line).strip() for line in (question_lines or lines)]
        )
        return [], {
            "rawCharacters": len(text),
            "warnings": [] if knowledge_points else ["未识别出题目或知识点"],
            "recognized": len(knowledge_points),
            "recognizedQuestions": 0,
            "recognizedKnowledgePoints": len(knowledge_points),
            "contentType": "knowledge",
            "knowledgePoints": knowledge_points,
        }

    blocks: list[tuple[list[str], str | None]] = []
    current: list[str] = []
    current_section: str | None = None
    seen_question = False
    for line in question_lines:
        heading = section_heading(line)
        if heading:
            if current:
                blocks.append((current, current_section))
                current = []
            current_section = heading
            continue
        if is_question_start(line) and current:
            blocks.append((current, current_section))
            current = [line]
            seen_question = True
        elif is_question_start(line):
            current = [line]
            seen_question = True
        elif seen_question:
            current.append(line)
        else:
            continue
    if current:
        blocks.append((current, current_section))

    questions: list[dict[str, Any]] = []
    parser_warnings: list[str] = []
    for index, (block, section) in enumerate(blocks, start=1):
        if not block:
            continue
        start_match = QUESTION_START.match(block[0])
        number = start_match.group(1) if start_match else str(index)
        raw_stem_source = question_stem(start_match) if start_match else block[0]
        question_images: list[str] = []
        for image_url in IMAGE_MARKER.findall(raw_stem_source):
            if image_url not in question_images:
                question_images.append(image_url)
        stem_source = IMAGE_MARKER.sub("", raw_stem_source).strip()
        stem = stem_source
        stem_inline_matches = list(INLINE_OPTION_START.finditer(stem_source))
        if len(stem_inline_matches) >= 2:
            stem = stem_source[:stem_inline_matches[0].start()].strip()
        options: list[dict[str, str]] = []
        question_lines: list[str] = []
        answer: list[str] = []
        explanation_lines: list[str] = []
        in_explanation = False
        for raw_line in ([raw_stem_source] + block[1:]):
            for image_url in IMAGE_MARKER.findall(raw_line):
                if image_url not in question_images:
                    question_images.append(image_url)
            line = IMAGE_MARKER.sub("", raw_line).strip()
            if not line:
                continue
            if ANSWER_LINE.match(line):
                answer = clean_answer(ANSWER_LINE.match(line).group(1))
                in_explanation = False
                continue
            explanation_match = EXPLANATION_LINE.match(line)
            if explanation_match:
                explanation_text = explanation_match.group(1).strip()
                if explanation_text:
                    explanation_lines.append(explanation_text)
                in_explanation = True
                continue
            if in_explanation:
                explanation_lines.append(line)
                continue
            inline_options = split_inline_options(line)
            option_match = OPTION_START.match(line)
            if len(inline_options) >= 2:
                options.extend(inline_options)
            elif option_match:
                options.append({"key": option_match.group(1).upper(), "text": option_match.group(2).strip()})
            elif line != stem_source:
                question_lines.append(line)

        question_text = TRAILING_MARKERS.sub("", " ".join([stem] + question_lines)).strip()
        matched_record = answer_matcher.take(number, section)
        if matched_record is not None:
            answer = matched_record.answer
            if matched_record.explanation:
                explanation_lines = [matched_record.explanation]
        elif number in answer_map:
            answer = answer_map[number]
            if explanation_map.get(number):
                explanation_lines = [explanation_map[number]]
        elif not answer:
            for line in block:
                if "正确答案" in line or "参考答案" in line:
                    answer = clean_answer(line.split(":", 1)[-1].split("：", 1)[-1])
                    break
        if not answer:
            parser_warnings.append(f"第 {number} 题未识别到答案")

        q_type = question_type(question_text, answer, options)
        if q_type == "判断题" and not options:
            options = [{"key": "对", "text": "正确"}, {"key": "错", "text": "错误"}]
        if q_type in {"单选题", "多选题"} and not options:
            parser_warnings.append(f"第 {number} 题未识别到选项")

        questions.append(
            {
                "id": str(uuid.uuid4()),
                "number": number,
                "type": q_type,
                "text": question_text,
                "options": options,
                "answer": answer,
                "explanation": " ".join(explanation_lines).strip()
                or "暂无解析，请在题库中补充。",
                "images": question_images,
                "source": "自动解析",
                "needsReview": not bool(answer) or (q_type in {"单选题", "多选题"} and not options),
            }
        )

    structured_questions = sum(1 for question in questions if question["options"] or question["answer"])
    if not structured_questions and text.strip():
        knowledge_points = extract_knowledge_points(
            [IMAGE_MARKER.sub("", line).strip() for line in (question_lines or lines)]
        )
        return [], {
            "rawCharacters": len(text),
            "warnings": [],
            "recognized": len(knowledge_points),
            "recognizedQuestions": 0,
            "recognizedKnowledgePoints": len(knowledge_points),
            "contentType": "knowledge",
            "knowledgePoints": knowledge_points,
        }

    return questions, {
        "rawCharacters": len(text),
        "warnings": parser_warnings,
        "recognized": len(questions),
        "recognizedQuestions": len(questions),
        "recognizedKnowledgePoints": 0,
        "contentType": "questions",
        "imageCount": sum(len(question.get("images", [])) for question in questions),
        "knowledgePoints": [],
    }


def split_inline_options(line: str) -> list[dict[str, str]]:
    """Split formats such as `A. one B. two C. three` into option records."""
    matches = list(INLINE_OPTION_START.finditer(line))
    if len(matches) < 2:
        return []
    options: list[dict[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(line)
        option_text = line[match.end():end].strip(" \t;；")
        if option_text:
            options.append({"key": match.group(1).upper(), "text": option_text})
    return options


def extract_knowledge_points(lines: list[str]) -> list[dict[str, Any]]:
    points: list[dict[str, Any]] = []
    seen: set[str] = set()
    for line in lines:
        value = re.sub(r"^\s*(?:[•·▪●○◦\-–—*]\s*)", "", line).strip()
        value = re.sub(r"^\s*(?:知识点|要点|重点)\s*[:：]\s*", "", value).strip()
        value = re.sub(r"^\s*(?:\d{1,4}\s*[\.．、\)）])\s*", "", value).strip()
        if not value or value in {"知识点", "要点", "重点", "目录"}:
            continue
        normalized = re.sub(r"\s+", "", value).casefold()
        if normalized in seen:
            continue
        seen.add(normalized)
        points.append(
            {
                "id": str(uuid.uuid4()),
                "text": value,
                "source": "自动整理",
            }
        )
    return points


def merge_knowledge_points(existing: list[dict[str, Any]], incoming: list[dict[str, Any]]) -> None:
    seen = {re.sub(r"\s+", "", item.get("text", "")).casefold() for item in existing}
    for item in incoming:
        key = re.sub(r"\s+", "", item.get("text", "")).casefold()
        if key and key not in seen:
            existing.append(item)
            seen.add(key)


def merge_questions(existing: list[dict[str, Any]], incoming: list[dict[str, Any]]) -> None:
    """Update an existing question when the same numbered stem is re-imported."""
    index = {
        (
            str(item.get("number", "")),
            re.sub(r"\s+", "", str(item.get("text", ""))).casefold(),
        ): item
        for item in existing
    }
    for question in incoming:
        key = (
            str(question.get("number", "")),
            re.sub(r"\s+", "", str(question.get("text", ""))).casefold(),
        )
        previous = index.get(key)
        if previous is None:
            existing.append(question)
            index[key] = question
            continue
        manually_corrected = previous.get("answerSource") == "manual"
        options_manually_corrected = previous.get("optionsSource") == "manual"
        explanation_manually_corrected = previous.get("explanationSource") == "manual"
        has_manual_revision = (
            manually_corrected
            or options_manually_corrected
            or explanation_manually_corrected
        )
        updates = {
            "images": question.get("images", previous.get("images", [])),
            "source": previous.get("source", "人工勘误") if has_manual_revision else question.get(
                "source", previous.get("source", "自动解析")
            ),
            "needsReview": False if manually_corrected else question.get(
                "needsReview", previous.get("needsReview", True)
            ),
        }
        if not manually_corrected:
            updates.update(
                {
                    "type": question.get("type", previous.get("type", "单选题")),
                    "answer": question.get("answer", previous.get("answer", [])),
                }
            )
        if not options_manually_corrected:
            updates["options"] = question.get("options", previous.get("options", []))
        if previous.get("explanationSource") != "manual":
            updates["explanation"] = question.get(
                "explanation",
                previous.get("explanation", ""),
            )
        previous.update(
            updates
        )


def module_summary(module: dict[str, Any]) -> dict[str, Any]:
    questions = module.get("questions", [])
    knowledge_points = module.get("knowledgePoints", [])
    pending = sum(1 for question in questions if question.get("needsReview"))
    types: dict[str, int] = {}
    for question in questions:
        q_type = question.get("type", "未分类")
        types[q_type] = types.get(q_type, 0) + 1
    return {
        "id": module["id"],
        "name": module["name"],
        "questionCount": len(questions),
        "knowledgeCount": len(knowledge_points),
        "contentCount": len(questions) + len(knowledge_points),
        "contentType": "mixed" if questions and knowledge_points else ("knowledge" if knowledge_points else "questions"),
        "pendingCount": pending,
        "types": types,
        "createdAt": module.get("createdAt"),
        "updatedAt": module.get("updatedAt"),
        "lastSource": module.get("lastSource"),
    }


def find_question(db: dict[str, Any], module_id: str, question_id: str) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    module = next((item for item in db["modules"] if item.get("id") == module_id), None)
    if not module:
        return None, None
    question = next((item for item in module.get("questions", []) if item.get("id") == question_id), None)
    return module, question


def practice_progress_key(module_id: str, question_id: str) -> str:
    return f"{module_id}::{question_id}"


def read_json_body(handler: BaseHTTPRequestHandler) -> dict[str, Any]:
    content_length = int(handler.headers.get("Content-Length", "0") or 0)
    if content_length <= 0:
        return {}
    try:
        return json.loads(handler.rfile.read(content_length).decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return {}


def json_response(handler: BaseHTTPRequestHandler, payload: Any, status: int = 200) -> None:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


class QuestionBankHandler(BaseHTTPRequestHandler):
    server_version = "QuestionBank/1.0"

    def log_message(self, format: str, *args: Any) -> None:
        print(f"[server] {self.address_string()} - {format % args}")

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/modules":
            db = load_db()
            json_response(self, {"modules": [module_summary(module) for module in db["modules"]]})
            return
        if parsed.path == "/api/practice-progress":
            db = load_db()
            module_id = parse_qs(parsed.query).get("moduleId", [""])[0]
            progress = db.get("practiceProgress", {})
            if module_id:
                prefix = f"{module_id}::"
                progress = {
                    key: value
                    for key, value in progress.items()
                    if key.startswith(prefix)
                }
            json_response(self, {"progress": progress})
            return
        if parsed.path.startswith("/api/modules/"):
            module_id = parsed.path.rsplit("/", 1)[-1]
            db = load_db()
            module = next((item for item in db["modules"] if item["id"] == module_id), None)
            if not module:
                json_response(self, {"error": "题库模块不存在"}, HTTPStatus.NOT_FOUND)
                return
            json_response(self, {"module": module})
            return
        if parsed.path == "/api/wrong-book":
            db = load_db()
            entries: list[dict[str, Any]] = []
            for entry in db.get("wrongBook", []):
                module, question = find_question(db, entry.get("moduleId", ""), entry.get("questionId", ""))
                if module and question:
                    entries.append(
                        {
                            "moduleId": module["id"],
                            "moduleName": module["name"],
                            "question": question,
                            "addedAt": entry.get("addedAt"),
                            "category": normalize_category(entry.get("category", "未分类")),
                            "note": entry.get("note", ""),
                            "manualAnswer": entry.get("manualAnswer", []),
                        }
                    )
            json_response(self, {"entries": entries})
            return
        self.serve_static(parsed.path)

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/wrong-book":
            payload = read_json_body(self)
            module_id = str(payload.get("moduleId", ""))
            question_id = str(payload.get("questionId", ""))
            db = load_db()
            module, question = find_question(db, module_id, question_id)
            if not module or not question:
                json_response(self, {"error": "题目不存在"}, HTTPStatus.NOT_FOUND)
                return
            existing = next(
                (
                    item
                    for item in db["wrongBook"]
                    if item.get("moduleId") == module_id and item.get("questionId") == question_id
                ),
                None,
            )
            category = normalize_category(payload.get("category", "未分类"))
            manual_answer = payload.get("manualAnswer", [])
            if not isinstance(manual_answer, list):
                manual_answer = []
            if existing:
                existing["category"] = category
                existing["note"] = str(payload.get("note", existing.get("note", ""))).strip()
                existing["manualAnswer"] = manual_answer
                save_db(db)
            else:
                db["wrongBook"].append(
                    {
                        "moduleId": module_id,
                        "questionId": question_id,
                        "addedAt": utc_now(),
                        "category": category,
                        "note": str(payload.get("note", "")).strip(),
                        "manualAnswer": manual_answer,
                    }
                )
                save_db(db)
            json_response(
                self,
                {"added": True, "moduleId": module_id, "questionId": question_id, "category": category},
                HTTPStatus.CREATED,
            )
            return
        if parsed.path != "/api/upload":
            json_response(self, {"error": "接口不存在"}, HTTPStatus.NOT_FOUND)
            return
        content_type = self.headers.get("Content-Type", "")
        content_length = int(self.headers.get("Content-Length", "0") or 0)
        if "multipart/form-data" not in content_type or content_length <= 0:
            json_response(self, {"error": "请通过文件上传方式提交"}, HTTPStatus.BAD_REQUEST)
            return
        multipart_body = self.rfile.read(content_length)
        uploads = read_uploaded_files(multipart_body, content_type)
        if not uploads:
            json_response(self, {"error": "没有找到上传文件"}, HTTPStatus.BAD_REQUEST)
            return
        try:
            question_upload, answer_upload = classify_uploads(uploads)
            all_uploads = [question_upload] + ([answer_upload] if answer_upload else [])
            for _, filename, content in all_uploads:
                extension = Path(filename).suffix.lower()
                if extension not in SUPPORTED_EXTENSIONS:
                    raise ValueError(
                        f"暂不支持 {extension or '该'} 文件，请上传 PDF、DOCX、PPTX、TXT 或 MD"
                    )
                if not content:
                    raise ValueError(f"文件「{filename}」内容为空")

            question_filename, question_content = question_upload[1], question_upload[2]
            question_text = read_file_text(question_filename, question_content, MEDIA_DIR)
            external_answer_data = None
            answer_filename = None
            answer_map: dict[str, list[str]] = {}
            if answer_upload:
                answer_filename, answer_content = answer_upload[1], answer_upload[2]
                answer_text = read_file_text(answer_filename, answer_content)
                external_answer_data = prepare_answer_data(answer_text)
                answer_map = external_answer_data[0]
                if not answer_map:
                    raise ValueError(f"未从答案文件「{answer_filename}」中识别出题号和答案")

            questions, parse_meta = parse_questions(question_text, external_answer_data)
            knowledge_points = parse_meta.get("knowledgePoints", [])
            if answer_upload:
                question_numbers = {question["number"] for question in questions}
                unmatched_numbers = sorted(set(answer_map) - question_numbers, key=lambda value: int(value))
                parse_meta["paired"] = True
                parse_meta["questionFile"] = question_filename
                parse_meta["answerFile"] = answer_filename
                parse_meta["answerCount"] = len(answer_map)
                if unmatched_numbers:
                    parse_meta.setdefault("warnings", []).append(
                        f"答案文件中有 {len(unmatched_numbers)} 个题号未在题目文件中找到："
                        + "、".join(unmatched_numbers[:20])
                    )
            else:
                parse_meta["paired"] = False
        except Exception as exc:
            json_response(self, {"error": f"解析失败：{exc}"}, HTTPStatus.BAD_REQUEST)
            return

        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        for _, filename, content in [question_upload] + ([answer_upload] if answer_upload else []):
            safe_filename = re.sub(r"[^0-9A-Za-z一-龥._-]+", "_", filename) or "upload"
            (UPLOAD_DIR / f"{timestamp}_{safe_filename}").write_bytes(content)

        module_name = normalize_name(Path(question_filename).stem)
        db = load_db()
        module = next(
            (
                item
                for item in db["modules"]
                if normalize_name(item.get("name", "")).casefold() == module_name.casefold()
            ),
            None,
        )
        now = utc_now()
        mode = "merged"
        if module:
            merge_questions(module.setdefault("questions", []), questions)
            merge_knowledge_points(module.setdefault("knowledgePoints", []), knowledge_points)
            module["updatedAt"] = now
            module["lastSource"] = (
                f"{question_filename} + {answer_filename}" if answer_filename else question_filename
            )
            module["imports"] = int(module.get("imports", 0)) + 1
        else:
            mode = "created"
            module = {
                "id": str(uuid.uuid4()),
                "name": module_name,
                "questions": questions,
                "knowledgePoints": knowledge_points,
                "createdAt": now,
                "updatedAt": now,
                "lastSource": (
                    f"{question_filename} + {answer_filename}" if answer_filename else question_filename
                ),
                "imports": 1,
            }
            db["modules"].insert(0, module)
        save_db(db)
        json_response(
            self,
            {
                "mode": mode,
                "filename": question_filename,
                "filenames": [item[1] for item in [question_upload] + ([answer_upload] if answer_upload else [])],
                "module": module,
                "parse": parse_meta,
            },
            HTTPStatus.CREATED,
        )

    def do_PUT(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/practice-progress":
            payload = read_json_body(self)
            module_id = str(payload.get("moduleId", ""))
            question_id = str(payload.get("questionId", ""))
            db = load_db()
            module, question = find_question(db, module_id, question_id)
            if not module or not question:
                json_response(self, {"error": "题目不存在"}, HTTPStatus.NOT_FOUND)
                return
            selected_answer = payload.get("lastAnswer", [])
            if not isinstance(selected_answer, list):
                selected_answer = []
            selected_answer = [str(item).strip() for item in selected_answer if str(item).strip()]
            correct = payload.get("lastCorrect")
            if not isinstance(correct, bool):
                correct = None
            try:
                elapsed = max(0, int(payload.get("lastElapsedSeconds", 0) or 0))
            except (TypeError, ValueError):
                elapsed = 0

            key = practice_progress_key(module_id, question_id)
            previous = db.setdefault("practiceProgress", {}).get(key, {})
            record = {
                "moduleId": module_id,
                "questionId": question_id,
                "lastAnswer": selected_answer,
                "lastCorrect": correct,
                "lastElapsedSeconds": elapsed,
                "totalElapsedSeconds": int(previous.get("totalElapsedSeconds", 0) or 0) + elapsed,
                "attempts": int(previous.get("attempts", 0) or 0) + 1,
                "lastAnsweredAt": utc_now(),
            }
            db["practiceProgress"][key] = record
            save_db(db)
            json_response(self, {"updated": True, "key": key, "progress": record})
            return
        if parsed.path.startswith("/api/modules/") and "/questions/" in parsed.path:
            path_parts = [unquote(part) for part in parsed.path.rstrip("/").split("/") if part]
            if len(path_parts) != 5 or path_parts[0:2] != ["api", "modules"] or path_parts[3] != "questions":
                json_response(self, {"error": "题目勘误地址无效"}, HTTPStatus.BAD_REQUEST)
                return
            module_id = path_parts[2]
            question_id = path_parts[4]
            payload = read_json_body(self)

            db = load_db()
            module, question = find_question(db, module_id, question_id)
            if not module or not question:
                json_response(self, {"error": "题目不存在"}, HTTPStatus.NOT_FOUND)
                return
            options = (
                normalize_manual_options(payload.get("options"))
                if "options" in payload
                else question.get("options", [])
            )
            answer = (
                normalize_manual_answer(payload.get("answer"), options)
                if "answer" in payload
                else question.get("answer", [])
            )
            question["options"] = options
            question["answer"] = answer
            question["type"] = question_type(
                str(question.get("text", "")),
                answer,
                options,
            )
            question["needsReview"] = not bool(answer) or (
                question["type"] in {"单选题", "多选题"} and not options
            )
            if "options" in payload:
                question["optionsSource"] = "manual"
                question["optionsUpdatedAt"] = utc_now()
            if "answer" in payload:
                question["answerSource"] = "manual"
                question["answerUpdatedAt"] = utc_now()
            question["source"] = "人工勘误"
            if "explanation" in payload:
                explanation = str(payload.get("explanation") or "").strip()
                question["explanation"] = explanation or "暂无解析，请在题库中补充。"
                question["explanationSource"] = "manual"
                question["explanationUpdatedAt"] = utc_now()
            module["updatedAt"] = utc_now()
            save_db(db)
            json_response(self, {"updated": True, "moduleId": module_id, "question": question})
            return
        if parsed.path != "/api/wrong-book":
            json_response(self, {"error": "接口不存在"}, HTTPStatus.NOT_FOUND)
            return
        payload = read_json_body(self)
        module_id = str(payload.get("moduleId", ""))
        question_id = str(payload.get("questionId", ""))
        db = load_db()
        entry = next(
            (
                item
                for item in db["wrongBook"]
                if item.get("moduleId") == module_id and item.get("questionId") == question_id
            ),
            None,
        )
        if not entry:
            json_response(self, {"error": "错题本中不存在这道题"}, HTTPStatus.NOT_FOUND)
            return
        if "category" in payload:
            entry["category"] = normalize_category(payload["category"])
        if "note" in payload:
            entry["note"] = str(payload.get("note") or "").strip()
        if "manualAnswer" in payload and isinstance(payload["manualAnswer"], list):
            entry["manualAnswer"] = payload["manualAnswer"]
        save_db(db)
        json_response(self, {"updated": True, "entry": entry})

    def do_DELETE(self) -> None:
        parsed = urlparse(self.path)
        request_path = parsed.path.rstrip("/")
        if request_path.startswith("/api/modules/") and "/questions/" in request_path:
            path_parts = [unquote(part) for part in request_path.split("/") if part]
            if len(path_parts) != 5 or path_parts[0:2] != ["api", "modules"] or path_parts[3] != "questions":
                json_response(self, {"error": "题目删除地址无效"}, HTTPStatus.BAD_REQUEST)
                return
            module_id = path_parts[2]
            question_id = path_parts[4]
            db = load_db()
            module = next((item for item in db["modules"] if item.get("id") == module_id), None)
            if not module:
                json_response(self, {"error": "题库模块不存在"}, HTTPStatus.NOT_FOUND)
                return
            questions = module.setdefault("questions", [])
            question = next((item for item in questions if item.get("id") == question_id), None)
            if not question:
                json_response(self, {"error": "题目不存在"}, HTTPStatus.NOT_FOUND)
                return
            module["questions"] = [item for item in questions if item.get("id") != question_id]
            db["wrongBook"] = [
                item
                for item in db.get("wrongBook", [])
                if not (item.get("moduleId") == module_id and item.get("questionId") == question_id)
            ]
            db["practiceProgress"].pop(practice_progress_key(module_id, question_id), None)
            module["updatedAt"] = utc_now()
            save_db(db)
            json_response(
                self,
                {
                    "removed": True,
                    "moduleId": module_id,
                    "questionId": question_id,
                    "remainingQuestions": len(module["questions"]),
                },
            )
            return
        if request_path.startswith("/api/modules/"):
            path_parts = [unquote(part) for part in request_path.split("/") if part]
            if len(path_parts) != 3 or path_parts[0:2] != ["api", "modules"]:
                json_response(self, {"error": "题库删除地址无效"}, HTTPStatus.BAD_REQUEST)
                return
            module_id = path_parts[2]
            db = load_db()
            module = next((item for item in db["modules"] if item.get("id") == module_id), None)
            if not module:
                json_response(self, {"error": "题库模块不存在"}, HTTPStatus.NOT_FOUND)
                return
            db["modules"] = [item for item in db["modules"] if item.get("id") != module_id]
            db["wrongBook"] = [
                item for item in db.get("wrongBook", []) if item.get("moduleId") != module_id
            ]
            prefix = f"{module_id}::"
            db["practiceProgress"] = {
                key: value
                for key, value in db.get("practiceProgress", {}).items()
                if not key.startswith(prefix)
            }
            save_db(db)
            json_response(
                self,
                {
                    "removed": True,
                    "moduleId": module_id,
                    "removedQuestions": len(module.get("questions", [])),
                    "removedKnowledgePoints": len(module.get("knowledgePoints", [])),
                },
            )
            return
        if parsed.path != "/api/wrong-book":
            json_response(self, {"error": "接口不存在"}, HTTPStatus.NOT_FOUND)
            return
        payload = read_json_body(self)
        module_id = str(payload.get("moduleId", ""))
        question_id = str(payload.get("questionId", ""))
        db = load_db()
        before = len(db["wrongBook"])
        db["wrongBook"] = [
            item
            for item in db["wrongBook"]
            if not (item.get("moduleId") == module_id and item.get("questionId") == question_id)
        ]
        if len(db["wrongBook"]) != before:
            save_db(db)
        json_response(self, {"removed": True})

    def serve_static(self, path: str) -> None:
        if path in {"", "/"}:
            path = "/index.html"
        relative = unquote(path.lstrip("/"))
        if path.startswith("/media/"):
            media_relative = unquote(path[len("/media/"):])
            file_path = (MEDIA_DIR / media_relative).resolve()
            allowed_root = MEDIA_DIR.resolve()
        else:
            file_path = (PUBLIC_DIR / relative).resolve()
            allowed_root = PUBLIC_DIR.resolve()
        if allowed_root not in file_path.parents and file_path != allowed_root:
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        if not file_path.is_file():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        content = file_path.read_bytes()
        content_type = mimetypes.guess_type(file_path.name)[0] or "application/octet-stream"
        if file_path.suffix == ".js":
            content_type = "application/javascript; charset=utf-8"
        elif file_path.suffix == ".css":
            content_type = "text/css; charset=utf-8"
        elif file_path.suffix == ".html":
            content_type = "text/html; charset=utf-8"
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(content)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(content)


def seed_sample_data() -> None:
    db = load_db()
    if db["modules"]:
        return
    now = utc_now()
    db["modules"] = [
        {
            "id": "demo-module",
            "name": "导入示例题库",
            "createdAt": now,
            "updatedAt": now,
            "lastSource": "导入示例题库.docx",
            "imports": 1,
            "questions": [
                {
                    "id": "demo-1",
                    "number": "1",
                    "type": "单选题",
                    "text": "下列哪一项最适合用来描述结构化题库的基本组成？",
                    "options": [
                        {"key": "A", "text": "题干、选项、答案与解析"},
                        {"key": "B", "text": "只有题目标题"},
                        {"key": "C", "text": "只有原始文件名"},
                        {"key": "D", "text": "只有分数记录"},
                    ],
                    "answer": ["A"],
                    "explanation": "刷题页面至少需要题干、可选答案、正确答案和解析，才能完成判分与复盘。",
                    "source": "示例数据",
                    "needsReview": False,
                },
                {
                    "id": "demo-2",
                    "number": "2",
                    "type": "判断题",
                    "text": "同名文件再次导入时，可以将新识别出的题目追加到原有题库模块中。",
                    "options": [{"key": "对", "text": "正确"}, {"key": "错", "text": "错误"}],
                    "answer": ["对"],
                    "explanation": "系统使用文件名（不含扩展名）作为模块名，同名模块会进入合并流程。",
                    "source": "示例数据",
                    "needsReview": False,
                },
                {
                    "id": "demo-3",
                    "number": "3",
                    "type": "多选题",
                    "text": "下面哪些格式可以被当前版本直接尝试解析？",
                    "options": [
                        {"key": "A", "text": "PDF"},
                        {"key": "B", "text": "DOCX"},
                        {"key": "C", "text": "PPTX"},
                        {"key": "D", "text": "MD"},
                    ],
                    "answer": ["A", "B", "C", "D"],
                    "explanation": "当前版本支持 PDF、DOCX、PPTX、TXT 和 MD；扫描型 PDF 若无文本层，仍需人工校正。",
                    "source": "示例数据",
                    "needsReview": False,
                },
            ],
        }
    ]
    save_db(db)


def main() -> None:
    ensure_storage()
    seed_sample_data()
    port = 8765
    server = ThreadingHTTPServer(("127.0.0.1", port), QuestionBankHandler)
    print(f"题库服务已启动：http://127.0.0.1:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
